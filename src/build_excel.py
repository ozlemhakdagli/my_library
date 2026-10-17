#!/usr/bin/env python3
"""data/raw/batch_*.jsonl -> kitaplar.xlsx + kitaplar.csv (PLAN.md Faz 4).

Kaynak JSON dosyalarına dokunmaz; her çalıştırmada çıktıları baştan üretir.
Usage: .venv/bin/python src/build_excel.py
"""
import csv
import difflib
import glob
import json
import os
import re
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_GLOB = os.path.join(ROOT, "data", "raw", "batch_*.jsonl")
XLSX_PATH = os.path.join(ROOT, "kitaplar.xlsx")
CSV_PATH = os.path.join(ROOT, "kitaplar.csv")
PHOTO_DIR = "<books-folder>"

UNREADABLE_TITLE = "[okunamadı]"
DUP_THRESHOLD = 0.88
CONF_RANK = {"high": 3, "medium": 2, "low": 1}
CONF_TR = {"high": "Yüksek", "medium": "Orta", "low": "Düşük"}
KIND_TR = {"cover": "Kapak", "spine": "Sırt"}
TR_ALPHABET = "abcçdefgğhıijklmnoöpqrsştuüvwxyz"
TR_RANK = {c: i for i, c in enumerate(TR_ALPHABET)}
_FOLD_ACCENTS = str.maketrans({"â": "a", "î": "i", "û": "u"})
MERGE_FIELDS = ("subtitle", "author", "publisher")
PARTIAL_MIN_LEN = 6  # kısmi okumada en az bu uzunlukta norm başlık (kısa başlıklar birleşmez)
CILT_RE = re.compile(r"\b[Cc]ilt\s+([IVXLCDM]+|\d+)\b")


# ---------------------------------------------------------------- saf fonksiyonlar

def norm(s):
    """Karşılaştırma anahtarı: İ->i, I->ı, küçük harf, ı->i, aksan/çengel at, harf-rakam dışını at."""
    s = s.replace("İ", "i").replace("I", "ı").lower().replace("ı", "i")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return "".join(c for c in s if c.isalnum())


def fields_compatible(a, b):
    """Biri boş VEYA norm'lar eşit VEYA biri diğerini içeriyor (yazar/yayınevi/alt başlık için)."""
    na, nb = norm(a), norm(b)
    if not na or not nb:
        return True
    return na == nb or na in nb or nb in na


def authors_compatible(a, b):
    return fields_compatible(a, b)


def publishers_compatible(a, b):
    return fields_compatible(a, b)


def subtitles_compatible(a, b):
    return fields_compatible(a, b)


def tr_sort_key(s):
    """Türkçe alfabe sırası. Baştaki noktalama yok sayılır; rakamlar en başta;
    harfler TR_ALPHABET'teki yerine gider (aksanlı yabancı harf -> temel harf)."""
    f = s.replace("İ", "i").replace("I", "ı").lower().translate(_FOLD_ACCENTS)
    f = re.sub(r"^[\W_]+", "", f)
    key = []
    for c in f:
        if c in TR_RANK:
            key.append((1, TR_RANK[c]))
        elif c.isdigit():
            key.append((0, c))
        elif c.isalnum():
            base = unicodedata.normalize("NFKD", c)[0]
            if base in TR_RANK:
                key.append((1, TR_RANK[base]))
            else:
                key.append((2, ord(base)))
        else:
            key.append((-1, ord(c)))  # iç boşluk/noktalama: kelime sınırı gibi, harflerden önce
    return tuple(key)


def _filled(entry):
    return sum(1 for k in ("subtitle", "author", "publisher", "note") if entry.get(k, "").strip())


def _base(group):
    return max(group, key=lambda e: (CONF_RANK.get(e["confidence"], 0), _filled(e)))


def merge_group(group):
    """Aynı kitabın kayıtlarını tek kayda indirger. Silme yok; üyeler 'members' içinde (temel ilk)."""
    base = _base(group)
    members = [base] + [e for e in group if e is not base]
    merged = dict(base)
    for field in MERGE_FIELDS:
        if not merged[field].strip():
            for other in members[1:]:
                if other[field].strip():
                    merged[field] = other[field]
                    break
    photos = []
    for e in members:
        if e["photo"] not in photos:
            photos.append(e["photo"])
    merged["photo"] = base["photo"]
    merged["photos"] = photos
    merged["members"] = members
    partial_notes = [f"Kısmi okuma birleşti ({e['photo']}): {e['title']}"
                     + (f" - {e['note']}" if e["note"] else "")
                     for e in members[1:] if e.get("partial")]
    if partial_notes:
        merged["note"] = " | ".join(x for x in [base["note"]] + partial_notes if x)
    return merged


def split_unreadable(entries):
    readable = [e for e in entries if e["title"] != UNREADABLE_TITLE]
    unreadable = [e for e in entries if e["title"] == UNREADABLE_TITLE]
    return readable, unreadable


def _same_book_fields(e, m):
    """Yazar, yayınevi ve alt başlık çelişmiyor mu?"""
    return (authors_compatible(e["author"], m["author"])
            and publishers_compatible(e["publisher"], m["publisher"])
            and subtitles_compatible(e["subtitle"], m["subtitle"]))


def group_entries(entries):
    """Aynı norm(title) + her üyeyle çelişmeyen yazar/yayınevi/alt başlık + farklı fotoğraf
    => aynı grup. Aynı fotoğraftaki kayıtlar asla birleşmez. Girdi sırası korunur."""
    groups = []  # her eleman: [norm_title, [entry, ...]]
    for e in entries:
        nt = norm(e["title"])
        for gkey, members in groups:
            if (gkey == nt
                    and all(m["photo"] != e["photo"] and _same_book_fields(e, m) for m in members)):
                members.append(e)
                break
        else:
            groups.append([nt, [e]])
    return [members for _, members in groups]


def partial_score(low_title, target_title):
    """Kısmi okuma eşleşme skoru: (alt_dize_mi, benzerlik) veya None (eşleşmez).
    Norm başlık en az PARTIAL_MIN_LEN uzunlukta, rakam dizileri aynı olmalı.
    Alt dize eşleşmesi (1, oran) her zaman bulanık eşleşmeden (0, oran) önce gelir.
    Bulanık eşleşme (benzerlik >= eşik) yalnızca iki başlığın kelime sayısı aynıysa geçerlidir
    (aksi halde 'Yaş Almak' ~ 'Yaşamak' gibi farklı kitapları birleştirir)."""
    nl, nt = norm(low_title), norm(target_title)
    if len(nl) < PARTIAL_MIN_LEN or not nt:
        return None
    if _digit_runs(low_title) != _digit_runs(target_title):
        return None
    ratio = difflib.SequenceMatcher(None, nl, nt, autojunk=False).ratio()
    if nl in nt:
        return (1, ratio)
    if ratio >= DUP_THRESHOLD and len(low_title.split()) == len(target_title.split()):
        return (0, ratio)
    return None


def absorb_partials(groups):
    """Düşük güvenli grupları (kısmi/bulanık okuma), başka fotoğraftaki düşük-olmayan bir grubun
    başlığıyla eşleşiyorsa o gruba ekler. Birden fazla eşit en iyi aday varsa birleştirmez.
    Eklenen kayıtlar 'partial' işaretlenir. Girdi/çıktı: grup listesi (liste of entry-listesi)."""
    groups = [list(g) for g in groups]
    lows = [g for g in groups if _base(g)["confidence"] == "low"]
    targets = [g for g in groups if _base(g)["confidence"] != "low"]
    absorbed_ids = set()
    for low in lows:
        low_base = _base(low)
        low_photos = {e["photo"] for e in low}
        cands = []
        for t in targets:
            if any(e["photo"] in low_photos for e in t):
                continue
            t_base = _base(t)
            score = partial_score(low_base["title"], t_base["title"])
            if score is None:
                continue
            if not (authors_compatible(low_base["author"], t_base["author"])
                    and publishers_compatible(low_base["publisher"], t_base["publisher"])):
                continue
            cands.append((score, t))
        if not cands:
            continue
        cands.sort(key=lambda c: c[0], reverse=True)
        if len(cands) > 1 and cands[0][0] == cands[1][0]:
            continue  # belirsiz: eşit aday
        target = cands[0][1]
        for e in low:
            e["partial"] = True
        target.extend(low)
        absorbed_ids.add(id(low))
    return [g for g in groups if id(g) not in absorbed_ids]


def _digit_runs(title):
    return re.findall(r"\d+", title)


def possible_duplicates(bases, threshold=DUP_THRESHOLD):
    """Birleşmemiş grup temel kayıtları arasında benzer başlıklar: [(a, b, oran)], oran azalan.
    Koşullar (hepsi): benzerlik >= eşik; yazar uyumlu; yayınevi çelişmiyor; başlıklardaki
    rakam dizileri aynı; iki kayıt farklı fotoğraflardan."""
    pairs = []
    keys = [norm(b["title"]) for b in bases]
    for i in range(len(bases)):
        for j in range(i + 1, len(bases)):
            a, b = bases[i], bases[j]
            ratio = difflib.SequenceMatcher(None, keys[i], keys[j], autojunk=False).ratio()
            if ratio < threshold:
                continue
            if not authors_compatible(a["author"], b["author"]):
                continue
            if not publishers_compatible(a["publisher"], b["publisher"]):
                continue
            if _digit_runs(a["title"]) != _digit_runs(b["title"]):
                continue
            if set(a["photos"]) & set(b["photos"]):
                continue
            pairs.append((a, b, ratio))
    pairs.sort(key=lambda p: p[2], reverse=True)
    return pairs


def volume_label(note):
    """Notta 'Cilt I' / 'Cilt II' / 'Cilt 3' geçiyorsa cilt işareti, yoksa ''."""
    m = CILT_RE.search(note or "")
    return m.group(1) if m else ""


def label_volumes(main):
    """Ana listede norm(title) aynı kalan satırlardan notunda cilt bilgisi olanların görünen başlığına
    ' (Cilt X)' ekler. Yalnızca gösterim; birleştirme anahtarını etkilemez."""
    counts = {}
    for m in main:
        k = norm(m["title"])
        counts[k] = counts.get(k, 0) + 1
    for m in main:
        if counts[norm(m["title"])] < 2:
            continue
        vol = volume_label(m["members"][0]["note"])
        if vol:
            m["title"] = f"{m['title']} (Cilt {vol})"


def check_consistency(raw, main, absorbed, unreadable):
    """Ham kayıt = ana liste kitap sayısı + birleşme sayısı + okunamadı."""
    if raw != main + absorbed + unreadable:
        raise ValueError(
            f"Tutarsızlık: ham={raw} != ana={main} + birleşme={absorbed} + okunamadı={unreadable}"
        )


def control_reasons(merged):
    reasons = []
    if merged["confidence"] == "low":
        reasons.append("Düşük güven")
    elif merged["confidence"] == "medium":
        reasons.append("Orta güven")
    if not merged["author"].strip():
        reasons.append("Yazar eksik")
    return reasons


# ---------------------------------------------------------------- okuma

def load_raw(raw_glob=RAW_GLOB):
    """Döner: (photos, entries). photos: fotoğraf başına kayıt; entries: kitap başına düz kayıt."""
    photos, entries = [], []
    for path in sorted(glob.glob(raw_glob)):
        with open(path, encoding="utf-8") as f:
            for line in f:
                if not line.strip():
                    continue
                rec = json.loads(line)
                photos.append({"photo": rec["photo"], "photo_note": rec.get("photo_note", ""),
                               "count": len(rec["books"])})
                for book in rec["books"]:
                    entries.append({
                        "photo": rec["photo"],
                        "photo_note": rec.get("photo_note", ""),
                        "title": book["title"].strip(),
                        "subtitle": book.get("subtitle", "").strip(),
                        "author": book.get("author", "").strip(),
                        "publisher": book.get("publisher", "").strip(),
                        "kind": book.get("kind", ""),
                        "confidence": book.get("confidence", ""),
                        "note": book.get("note", "").strip(),
                    })
    return photos, entries


# ---------------------------------------------------------------- tablo üretimi

def build_tables(photos, entries):
    readable, unreadable = split_unreadable(entries)
    groups = absorb_partials(group_entries(readable))
    merged = [merge_group(g) for g in groups]
    merged.sort(key=lambda m: (tr_sort_key(m["title"]), tr_sort_key(m["subtitle"]), norm(m["author"])))
    absorbed = sum(len(m["members"]) - 1 for m in merged)
    check_consistency(len(entries), len(merged), absorbed, len(unreadable))

    duplicates = possible_duplicates(merged)
    label_volumes(merged)

    control = []  # (sıralama anahtarı, satır sözlüğü)
    for m in merged:
        reasons = control_reasons(m)
        if reasons:
            control.append(((0, tr_sort_key(m["title"])), {
                "title": m["title"], "subtitle": m["subtitle"], "author": m["author"],
                "publisher": m["publisher"], "kind": m["kind"], "confidence": m["confidence"],
                "photos": m["photos"], "reason": ", ".join(reasons), "note": m["note"],
            }))
    for e in unreadable:
        note = " | ".join(x for x in (e["photo_note"], e["note"]) if x)
        control.append(((1, e["photo"]), {
            "title": UNREADABLE_TITLE, "subtitle": "", "author": "", "publisher": "",
            "kind": e["kind"], "confidence": "low", "photos": [e["photo"]],
            "reason": "Okunamadı", "note": note,
        }))
    control.sort(key=lambda c: c[0])
    control_rows = [c[1] for c in control]

    return {
        "main": merged,
        "control": control_rows,
        "duplicates": duplicates,
        "photos": sorted(photos, key=lambda p: p["photo"]),
        "stats": {
            "photos": len(photos),
            "raw": len(entries),
            "main": len(merged),
            "groups_merged": sum(1 for m in merged if len(m["members"]) > 1),
            "absorbed": absorbed,
            "unreadable": len(unreadable),
            "control": len(control_rows),
            "publisher_empty": sum(1 for m in merged if not m["publisher"].strip()),
            "author_empty": sum(1 for m in merged if not m["author"].strip()),
            "multi_book_photos": sum(1 for p in photos if p["count"] > 1),
            "duplicates": len(duplicates),
        },
    }


# ---------------------------------------------------------------- yazma

def _photo_link(cell, name):
    cell.value = name
    cell.hyperlink = f"file://{os.path.join(PHOTO_DIR, name)}"
    cell.font = Font(color="0563C1", underline="single")


def _write_sheet(ws, headers, rows, widths, link_col=None):
    """rows: list of list; link_col: 0-based sütun indeksi, hücre köprüsü (fotoğraf adı)."""
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r in rows:
        ws.append(r)
        if link_col is not None and r[link_col]:
            _photo_link(ws.cell(row=ws.max_row, column=link_col + 1), r[link_col])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"


KITAP_HEADERS = ["No", "Başlık", "Alt Başlık", "Yazar", "Yayınevi", "Tür", "Güven",
                 "Fotoğraf Sayısı", "Fotoğraflar", "Not", "Kontrol"]
KITAP_WIDTHS = [6, 42, 48, 30, 20, 10, 10, 16, 40, 60, 10]


def _main_rows(main):
    rows = []
    for i, m in enumerate(main, start=1):
        rows.append([
            i, m["title"], m["subtitle"], m["author"], m["publisher"],
            KIND_TR.get(m["kind"], m["kind"]), CONF_TR.get(m["confidence"], m["confidence"]),
            len(m["photos"]), ", ".join(m["photos"]), m["note"],
            "Evet" if control_reasons(m) else "",
        ])
    return rows


def _write_workbook(tables, path):
    wb = Workbook()
    stats = tables["stats"]

    ws = wb.active
    ws.title = "Kitaplar"
    rows = _main_rows(tables["main"])
    _write_sheet(ws, KITAP_HEADERS, rows, KITAP_WIDTHS)
    for r in range(2, ws.max_row + 1):
        _photo_link(ws.cell(row=r, column=9), ws.cell(row=r, column=9).value.split(",")[0].strip())

    ws = wb.create_sheet("Kontrol")
    crow = []
    for i, c in enumerate(tables["control"], start=1):
        crow.append([i, c["title"], c["subtitle"], c["author"], c["publisher"],
                     KIND_TR.get(c["kind"], c["kind"]), CONF_TR.get(c["confidence"], c["confidence"]),
                     ", ".join(c["photos"]), c["reason"], c["note"]])
    _write_sheet(ws, ["No", "Başlık", "Alt Başlık", "Yazar", "Yayınevi", "Tür", "Güven",
                      "Fotoğraflar", "Neden", "Not"], crow,
                 [6, 42, 48, 30, 20, 10, 10, 40, 30, 60], link_col=7)

    ws = wb.create_sheet("Tekrarlar")
    ws.append(["Aynı kitap birden fazla fotoğrafta görünüyor; iki kopya sizde varsa Fotoğraf Sayısı'na bakın."])
    ws["A1"].font = Font(italic=True)
    t_headers = ["Bölüm", "Grup", "Rol", "Başlık", "Alt Başlık", "Yazar", "Yayınevi", "Tür",
                 "Güven", "Fotoğraf", "Grup Fotoğraf Sayısı", "Benzerlik"]
    trows = []
    gno = 0
    for m in tables["main"]:
        if len(m["members"]) < 2:
            continue
        gno += 1
        section = ("Birleşen (kısmi okuma)" if any(e.get("partial") for e in m["members"])
                   else "Birleşen")
        for k, e in enumerate(m["members"]):
            trows.append([section, f"G{gno}", "Temel" if k == 0 else "Birleşen", e["title"],
                          e["subtitle"], e["author"], e["publisher"], KIND_TR.get(e["kind"], e["kind"]),
                          CONF_TR.get(e["confidence"], e["confidence"]), e["photo"], len(m["photos"]), ""])
    for pno, (a, b, ratio) in enumerate(tables["duplicates"], start=1):
        for role, e in (("A", a), ("B", b)):
            trows.append(["Olası tekrar (birleştirilmedi)", f"P{pno}", role, e["title"],
                          e["subtitle"], e["author"], e["publisher"], KIND_TR.get(e["kind"], e["kind"]),
                          CONF_TR.get(e["confidence"], e["confidence"]), e["photos"][0] if e.get("photos") else e["photo"],
                          "", f"{ratio:.2f}" if role == "A" else ""])
    ws.append(t_headers)
    hdr = ws.max_row
    for c in range(1, len(t_headers) + 1):
        ws.cell(row=hdr, column=c).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = [30, 7, 10, 42, 48, 30, 20, 10, 10, 20, 20, 12][c - 1]
    for r in trows:
        ws.append(r)
        _photo_link(ws.cell(row=ws.max_row, column=10), r[9])
    ws.freeze_panes = f"A{hdr + 1}"
    ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(t_headers))}{ws.max_row}"

    ws = wb.create_sheet("Fotoğraf Özeti")
    prow = [[p["photo"], p["count"], p["photo_note"]] for p in tables["photos"]]
    _write_sheet(ws, ["Fotoğraf", "Kitap Sayısı", "Not"], prow, [24, 14, 90], link_col=0)

    ws = wb.create_sheet("Özet")
    summary = [
        ("Toplam foto", stats["photos"]),
        ("Toplam ham kayıt", stats["raw"]),
        ("Ana listedeki benzersiz kitap sayısı", stats["main"]),
        ("Birleşen grup sayısı", stats["groups_merged"]),
        ("Birleşme sayısı (birleşen ek kayıt)", stats["absorbed"]),
        ("Okunamadı sayısı", stats["unreadable"]),
        ("Kontrol'e giden kitap sayısı", stats["control"]),
        ("Yayınevi boş sayısı (ana liste)", stats["publisher_empty"]),
        ("Yazar boş sayısı (ana liste)", stats["author_empty"]),
        ("Çok kitaplı foto sayısı", stats["multi_book_photos"]),
        ("Olası tekrar sayısı (birleştirilmedi)", stats["duplicates"]),
        ("Tutarlılık (ham = ana + birleşme + okunamadı)", "Tamam"),
    ]
    _write_sheet(ws, ["Rubrik", "Değer"], [list(r) for r in summary], [46, 14])

    wb.save(path)


def _write_csv(tables, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(KITAP_HEADERS)
        w.writerows(_main_rows(tables["main"]))


def main():
    photos, entries = load_raw()
    tables = build_tables(photos, entries)
    _write_workbook(tables, XLSX_PATH)
    _write_csv(tables, CSV_PATH)
    s = tables["stats"]
    print(f"OK: {s['photos']} foto, {s['raw']} ham kayıt -> {s['main']} ana kitap, "
          f"{s['groups_merged']} birleşen grup ({s['absorbed']} ek kayıt), "
          f"{s['unreadable']} okunamadı, {s['control']} kontrol, {s['duplicates']} olası tekrar")
    print(f"Yazıldı: {XLSX_PATH}\n         {CSV_PATH}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
